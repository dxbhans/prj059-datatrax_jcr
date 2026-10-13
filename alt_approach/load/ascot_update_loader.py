"""Load and locally cache ASCOT going-stick, report, and calendar data."""

from __future__ import annotations

import logging
import os
import re
from pathlib import Path

import duckdb
import pandas as pd
import pymysql
from openpyxl import load_workbook

LOGGER = logging.getLogger(__name__)

DB_HOST = "tthubdb-cluster-1.cluster-che9ka0taxng.eu-west-1.rds.amazonaws.com"
DB_USER = "hubview"
DB_PASSWORD = "v13whub_"
DB_NAME = "tthub"
CALENDAR_ROOT = Path(__file__).resolve().parents[1] / "resources" / "calendars"

START_DATE = "2021-01-01"
END_DATE = "2027-01-01"
END_DATE_INCLUSIVE = "2026-12-31"
ASCOT_COURSES = ("Ascot",)

# The API user supplies automated Official Going values and must remain part of
# declaration/race-day evidence. The legacy display name is the only source
# excluded from that evidence.
AUTOMATED_OFFICIAL_GOING_USERS = {"apiuser.bhagoing@turftrax.co.uk"}
EXCLUDED_GOING_REPORT_USERS = {"tt auto updater"}


def configure_logging(level: int = logging.INFO) -> None:
    logging.basicConfig(level=level, format="%(asctime)s | %(levelname)s | %(message)s", force=True)
    LOGGER.setLevel(level)


def _database_connection() -> pymysql.connections.Connection:
    """Open the read-only Turftrax connection."""
    LOGGER.info("Connecting to the Turftrax database.")
    return pymysql.connect(
        host=os.getenv("TURFTRAX_DB_HOST", DB_HOST),
        user=os.getenv("TURFTRAX_DB_USER", DB_USER),
        password=os.getenv("TURFTRAX_DB_PASSWORD", DB_PASSWORD),
        database=os.getenv("TURFTRAX_DB_NAME", DB_NAME),
        cursorclass=pymysql.cursors.DictCursor,
    )


def _read_sql(sql_path: Path) -> str:
    return sql_path.read_text(encoding="utf-8")


def _fetch_query(sql_path: Path) -> pd.DataFrame:
    """Execute a SQL file and return its result as a dataframe."""
    LOGGER.info("Downloading %s.", sql_path.name)
    with _database_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute(_read_sql(sql_path))
            rows = cursor.fetchall()
            columns = [column[0] for column in cursor.description]
    dataframe = pd.DataFrame(rows, columns=columns)
    LOGGER.info("Downloaded %s rows from %s.", len(dataframe), sql_path.name)
    return dataframe


def _sync_duckdb(dataframe: pd.DataFrame, table_name: str, database_path: Path) -> None:
    """Store a queryable copy of a parquet-backed cache in DuckDB."""
    database_path.parent.mkdir(parents=True, exist_ok=True)
    with duckdb.connect(str(database_path)) as connection:
        connection.register("cache_frame", dataframe)
        connection.execute(f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM cache_frame')
        connection.unregister("cache_frame")
    LOGGER.info("Updated DuckDB table %s.", table_name)


def cached_query(
    sql_path: Path,
    cache_path: Path,
    table_name: str,
    duckdb_path: Path,
    refresh: bool = False,
) -> pd.DataFrame:
    """Read a parquet cache or download it once, then maintain DuckDB."""
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    if refresh:
        LOGGER.info("Refreshing cache for %s from %s.", table_name, sql_path.name)
    if cache_path.exists() and not refresh:
        LOGGER.info("Loading cached parquet: %s.", cache_path.name)
        dataframe = pd.read_parquet(cache_path)
    else:
        dataframe = _fetch_query(sql_path)
        dataframe.to_parquet(cache_path, index=False, compression="zstd")
        LOGGER.info("Saved parquet cache: %s.", cache_path.name)

    _sync_duckdb(dataframe, table_name, duckdb_path)
    return dataframe


def load_database_sources(load_dir: Path, data_dir: Path, refresh: bool = False) -> dict[str, pd.DataFrame]:
    """Load ASCOT readings, reports, calendars, and Going groups."""
    duckdb_path = data_dir / "ascot_2026.duckdb"
    sources = {
        "stick_readings": ("001_stick_readings.sql", "ascot_stick_readings.parquet"),
        "going_reports": ("002_going_reports.sql", "ascot_going_reports.parquet"),
        "calendar_database": ("003_calendar.sql", "ascot_calendar_database.parquet"),
        "going_groups": ("004_going_groups.sql", "ascot_going_groups.parquet"),
        "maps": ("005_maps.sql", "ascot_maps.parquet"),
        "map_waypoints": ("006_map_waypoints.sql", "ascot_map_waypoints.parquet"),
        "map_sections": ("007_map_sections.sql", "ascot_map_sections.parquet"),
        "map_zones": ("008_map_zones.sql", "ascot_map_zones.parquet"),
    }
    loaded_sources = {}
    for name, (sql_name, cache_name) in sources.items():
        sql_path = load_dir / sql_name
        LOGGER.info("Loading source %s (%s).", name, "refresh" if refresh else "cache-or-download")
        cache_path = data_dir / cache_name
        dataframe = cached_query(sql_path, cache_path, name, duckdb_path, refresh)
        if name == "going_reports" and "stick_user" not in dataframe.columns:
            LOGGER.info("Going-report cache predates stick-user data; refreshing it once.")
            dataframe = cached_query(
                sql_path, cache_path, name, duckdb_path, refresh=True
            )
        LOGGER.info("Loaded source %s: %s rows.", name, len(dataframe))
        loaded_sources[name] = dataframe
    return loaded_sources


def _normalise_course(series: pd.Series) -> pd.Series:
    normalised = series.astype(str).str.strip().str.replace(r"\s+", " ", regex=True).str.title()
    return normalised.replace({"Yarmouth": "Great Yarmouth"})


def _normalise_text(series: pd.Series) -> pd.Series:
    return series.astype("string").fillna("").str.strip().str.replace(r"\s+", " ", regex=True)


def _normalise_excel_for_parquet(dataframe: pd.DataFrame) -> pd.DataFrame:
    """Coerce mixed Excel object columns to stable Parquet-friendly types."""
    dataframe = dataframe.copy()
    for column in dataframe.select_dtypes(include=["object"]).columns:
        dataframe[column] = _normalise_text(dataframe[column])
    return dataframe


def load_excel_calendar(data_dir: Path, refresh: bool = False) -> pd.DataFrame:
    """Read ASCOT meeting rows from BHA Excel workbooks and cache them locally."""
    cache_path = data_dir / "ascot_calendar_excel.parquet"
    duckdb_path = data_dir / "ascot_2026.duckdb"
    if cache_path.exists() and not refresh:
        LOGGER.info("Loading cached parquet: %s.", cache_path.name)
        dataframe = pd.read_parquet(cache_path)
        dataframe["meeting_date"] = pd.to_datetime(dataframe["meeting_date"]).dt.normalize()
        dataframe["course"] = _normalise_course(dataframe["course"])
        if "excel_race_code" not in dataframe:
            source_code = "code" if "code" in dataframe else "meeting_code"
            dataframe["excel_race_code"] = dataframe[source_code].astype("string").replace(
                {"Both": "Mixed"}
            )
        _sync_duckdb(dataframe, "calendar_excel", duckdb_path)
        return dataframe

    calendar_root = Path(os.getenv("ASCOT_CALENDAR_ROOT", str(CALENDAR_ROOT)))
    if refresh:
        LOGGER.info("Refreshing Excel calendar cache from %s.", calendar_root)
    if not calendar_root.exists():
        raise FileNotFoundError(
            f"BHA calendar root does not exist: {calendar_root}. Set ASCOT_CALENDAR_ROOT."
        )

    workbooks = [
        calendar_root / f"BHA MASTER {year}_Fixture_List.xlsx"
        for year in range(2018, 2027)
    ]
    missing_workbooks = [path for path in workbooks if not path.exists()]
    if missing_workbooks:
        missing_list = "\n".join(str(path) for path in missing_workbooks)
        raise FileNotFoundError(f"Missing expected BHA Excel workbook(s):\n{missing_list}")

    LOGGER.info("Reading %s BHA Excel workbooks.", len(workbooks))
    frames = []
    for workbook in workbooks:
        LOGGER.info("Reading BHA Excel workbook: %s.", workbook)
        print(f"Reading BHA Excel workbook: {workbook}", flush=True)
        try:
            frame = pd.read_excel(workbook, sheet_name="List", engine="openpyxl", parse_dates=["Date"])
        except ValueError:
            try:
                workbook_sheets = load_workbook(workbook, read_only=True, data_only=True).sheetnames
            except Exception as sheet_error:
                workbook_sheets = [f"could not inspect sheet names: {sheet_error}"]
            LOGGER.exception(
                "Failed reading sheet 'List' from BHA Excel workbook: %s. Available sheets: %s",
                workbook,
                workbook_sheets,
            )
            print(
                f"Failed reading sheet 'List' from BHA Excel workbook: {workbook}. "
                f"Available sheets: {workbook_sheets}",
                flush=True,
            )
            raise
        frame.columns = frame.columns.str.strip().str.lower()
        required = {"date", "course"}
        if not required.issubset(frame.columns):
            LOGGER.warning("Skipping %s: missing %s.", workbook.name, required - set(frame.columns))
            continue
        frame["source_file"] = workbook.name
        frames.append(frame)

    if not frames:
        raise ValueError("No readable BHA calendar rows were found.")

    dataframe = pd.concat(frames, ignore_index=True)
    dataframe["meeting_date"] = pd.to_datetime(dataframe["date"]).dt.normalize()
    dataframe["course"] = _normalise_course(dataframe["course"])
    dataframe = dataframe[
        dataframe["course"].isin(ASCOT_COURSES)
        & dataframe["meeting_date"].between(START_DATE, END_DATE_INCLUSIVE)
    ].copy()
    dataframe["excel_race_code"] = dataframe.get(
        "code", pd.Series(pd.NA, index=dataframe.index)
    ).astype("string").replace({"Both": "Mixed"})
    dataframe["meeting_code"] = dataframe["excel_race_code"]
    dataframe = dataframe.sort_values(["meeting_date", "course"]).reset_index(drop=True)
    dataframe = _normalise_excel_for_parquet(dataframe)
    dataframe.to_parquet(cache_path, index=False, compression="zstd")
    LOGGER.info("Saved %s ASCOT Excel calendar rows to %s.", len(dataframe), cache_path.name)
    _sync_duckdb(dataframe, "calendar_excel", duckdb_path)
    return dataframe


def prepare_sources(sources: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Apply stable types and common date/course fields used by the notebook."""
    stick_readings = sources["stick_readings"].copy()
    stick_readings["sr_created_at"] = pd.to_datetime(stick_readings["sr_created_at"])
    stick_readings["date"] = stick_readings["sr_created_at"].dt.normalize()
    stick_readings["gr_race_date"] = pd.to_datetime(stick_readings["gr_race_date"], errors="coerce").dt.normalize()
    stick_readings["course"] = _normalise_course(stick_readings["course"])
    for column in ("sr_index", "sr_penetration", "sr_shear"):
        stick_readings[column] = pd.to_numeric(stick_readings[column], errors="coerce")
    stick_readings["data_type"] = "data point"
    stick_readings.loc[(stick_readings["sr_index"] == 0) | (stick_readings["sr_shear"] == 15), "data_type"] = "error"

    going_reports = sources["going_reports"].copy()
    going_reports["created_at_gr"] = pd.to_datetime(going_reports["created_at_gr"])
    going_reports["date"] = going_reports["created_at_gr"].dt.normalize()
    going_reports["race_date"] = pd.to_datetime(going_reports["race_date"], errors="coerce").dt.normalize()
    going_reports["course"] = _normalise_course(going_reports["course"])

    calendar_database = sources["calendar_database"].copy()
    calendar_database["meeting_date"] = pd.to_datetime(calendar_database["meeting_date"]).dt.normalize()
    calendar_database["course"] = _normalise_course(calendar_database["course"])
    calendar_database["database_meeting_code"] = calendar_database["meeting_code"].astype("string")

    going_groups = sources["going_groups"].copy()
    going_groups["course"] = _normalise_course(going_groups["course"])
    for column in ("going_min", "going_max"):
        going_groups[column] = pd.to_numeric(going_groups[column], errors="coerce")

    maps = sources["maps"].copy()
    maps["course"] = _normalise_course(maps["course"])
    map_waypoints = sources["map_waypoints"].copy()
    map_waypoints["course"] = _normalise_course(map_waypoints["course"])
    map_waypoints["waypoint"] = pd.to_numeric(map_waypoints["waypoint"], errors="coerce").astype("Int64")
    map_sections = sources["map_sections"].copy()
    map_sections["course"] = _normalise_course(map_sections["course"])
    map_zones = sources["map_zones"].copy()
    map_zones["course"] = _normalise_course(map_zones["course"])
    return {
        "stick_readings": stick_readings,
        "going_reports": going_reports,
        "calendar_database": calendar_database,
        "going_groups": going_groups,
        "maps": maps,
        "map_waypoints": map_waypoints,
        "map_sections": map_sections,
        "map_zones": map_zones,
    }


def calendar_discrepancies(
    calendar_database: pd.DataFrame, calendar_excel: pd.DataFrame
) -> tuple[bool, pd.DataFrame]:
    """Compare ASCOT fixtures by date and course, excluding source-specific codes."""
    reconciliation = calendar_reconciliation(calendar_database, calendar_excel)
    discrepancies = reconciliation.loc[
        reconciliation["reconciliation_status"].ne("matched")
    ].reset_index(drop=True)
    identical = discrepancies.empty
    LOGGER.info("Calendar comparison complete: %s.", "identical" if identical else f"{len(discrepancies)} discrepancies")
    return identical, discrepancies


def _joined_unique_values(series: pd.Series) -> str:
    values = sorted(_normalise_text(series).loc[lambda value: value.ne("")].unique())
    return " | ".join(values)


def calendar_reconciliation(
    calendar_database: pd.DataFrame, calendar_excel: pd.DataFrame
) -> pd.DataFrame:
    """Reconcile calendars on date and course while preserving source code semantics."""
    key_columns = ["meeting_date", "course"]
    database = calendar_database.copy()
    excel = calendar_excel.copy()
    for frame in (database, excel):
        frame["meeting_date"] = pd.to_datetime(frame["meeting_date"]).dt.normalize()
        frame["course"] = _normalise_course(frame["course"])

    database_code = "database_meeting_code" if "database_meeting_code" in database else "meeting_code"
    excel_code = "excel_race_code" if "excel_race_code" in excel else "meeting_code"
    database_summary = (
        database.groupby(key_columns, dropna=False)
        .agg(
            database_record_count=(database_code, "size"),
            database_meeting_codes=(database_code, _joined_unique_values),
        )
        .reset_index()
    )
    excel_summary = (
        excel.groupby(key_columns, dropna=False)
        .agg(
            excel_record_count=(excel_code, "size"),
            excel_race_codes=(excel_code, _joined_unique_values),
        )
        .reset_index()
    )
    result = database_summary.merge(excel_summary, on=key_columns, how="outer")
    result[["database_record_count", "excel_record_count"]] = result[
        ["database_record_count", "excel_record_count"]
    ].fillna(0).astype("int64")
    result[["database_meeting_codes", "excel_race_codes"]] = result[
        ["database_meeting_codes", "excel_race_codes"]
    ].fillna("")
    result["reconciliation_status"] = "matched"
    result.loc[result["excel_record_count"].eq(0), "reconciliation_status"] = "database_only"
    result.loc[result["database_record_count"].eq(0), "reconciliation_status"] = "excel_only"
    result.loc[
        result["database_record_count"].gt(1) & result["excel_record_count"].gt(0),
        "reconciliation_status",
    ] = "ambiguous_multiple_database_records"
    result.loc[
        result["excel_record_count"].gt(1) & result["database_record_count"].gt(0),
        "reconciliation_status",
    ] = "ambiguous_multiple_excel_records"
    result = result.sort_values(key_columns).reset_index(drop=True)
    result["declaration_date"] = result["meeting_date"] - pd.Timedelta(days=2)
    LOGGER.info("Calendar reconciliation complete: %s canonical fixtures.", len(result))
    return result


def calendar_duplicate_records(
    calendar_database: pd.DataFrame, calendar_excel: pd.DataFrame
) -> pd.DataFrame:
    """Return source records belonging to duplicated date-course fixture keys."""
    key_columns = ["meeting_date", "course"]
    results = []
    for source, frame, code_column in (
        ("database", calendar_database, "database_meeting_code"),
        ("excel", calendar_excel, "excel_race_code"),
    ):
        records = frame.copy()
        records["meeting_date"] = pd.to_datetime(records["meeting_date"]).dt.normalize()
        records["course"] = _normalise_course(records["course"])
        if code_column not in records:
            records[code_column] = records["meeting_code"].astype("string")
        duplicate_keys = records.duplicated(key_columns, keep=False)
        if duplicate_keys.any():
            detail_columns = key_columns + [code_column]
            detail_columns += [
                column
                for column in ("calendar_id", "note", "source_file", "created_at", "updated_at")
                if column in records
            ]
            detail = records.loc[duplicate_keys, detail_columns].copy()
            detail.insert(0, "source", source)
            detail = detail.rename(columns={code_column: "source_code"})
            results.append(detail)
    if not results:
        return pd.DataFrame(columns=["source"] + key_columns + ["source_code"])
    return pd.concat(results, ignore_index=True).sort_values(
        key_columns + ["source", "source_code"]
    ).reset_index(drop=True)


def calendar_field_discrepancies(
    calendar_database: pd.DataFrame, calendar_excel: pd.DataFrame
) -> pd.DataFrame:
    """Compare common fields for fixtures present in both calendar sources."""
    key_columns = ["meeting_date", "course"]
    database = calendar_database.copy()
    excel = calendar_excel.copy()
    for frame in (database, excel):
        frame["meeting_date"] = pd.to_datetime(frame["meeting_date"]).dt.normalize()
        frame["course"] = _normalise_course(frame["course"])

    database = database.loc[~database.duplicated(key_columns, keep=False)].copy()
    excel = excel.loc[~excel.duplicated(key_columns, keep=False)].copy()
    comparable_columns = sorted(
        (set(database.columns) & set(excel.columns))
        - set(key_columns)
        - {
            "date", "source_file", "created_at", "updated_at", "meeting_code",
            "database_meeting_code", "excel_race_code", "code",
        }
    )
    if not comparable_columns:
        LOGGER.info("Calendar field comparison skipped: no non-key common fields.")
        return pd.DataFrame(columns=key_columns + ["field", "database_value", "excel_value"])

    matched = database[key_columns + comparable_columns].merge(
        excel[key_columns + comparable_columns],
        on=key_columns,
        how="inner",
        suffixes=("_database", "_excel"),
    )
    differences = []
    for column in comparable_columns:
        left = _normalise_text(matched[f"{column}_database"])
        right = _normalise_text(matched[f"{column}_excel"])
        changed = left.ne(right)
        if changed.any():
            rows = matched.loc[changed, key_columns].copy()
            rows["field"] = column
            rows["database_value"] = left.loc[changed].to_numpy()
            rows["excel_value"] = right.loc[changed].to_numpy()
            differences.append(rows)

    if not differences:
        result = pd.DataFrame(columns=key_columns + ["field", "database_value", "excel_value"])
    else:
        result = pd.concat(differences, ignore_index=True).sort_values(
            key_columns + ["field"]
        ).reset_index(drop=True)
    LOGGER.info("Calendar field comparison complete: %s differences.", len(result))
    return result


def validate_going_groups(going_groups: pd.DataFrame) -> pd.DataFrame:
    """Audit map-specific Going bands for missing bounds, gaps, and overlaps."""
    key_columns = ["course", "map_id", "map_label", "goings_group_id", "goings_group"]
    bands = going_groups.drop_duplicates(
        ["map_id", "goings_group_id", "going_id"]
    ).copy()
    results = []
    for key, group in bands.groupby(key_columns, dropna=False):
        ordered = group.sort_values(["going_min", "going_max", "going_order"])
        inactive = _normalise_text(ordered["going"]).str.casefold().str.replace(" ", "").eq(
            "notinuse"
        )
        active = ordered.loc[~inactive]
        valid = active.loc[
            active["going_min"].notna()
            & active["going_max"].notna()
            & active["going_min"].le(active["going_max"])
        ]
        previous_max = valid["going_max"].shift()
        row = dict(zip(key_columns, key if isinstance(key, tuple) else (key,)))
        row.update(
            band_count=len(ordered),
            inactive_band_count=int(inactive.sum()),
            missing_bound_count=int(
                (active["going_min"].isna() | active["going_max"].isna()).sum()
            ),
            invalid_bound_count=int(active["going_min"].gt(active["going_max"]).sum()),
            overlap_count=int(valid["going_min"].lt(previous_max).sum()),
            gap_count=int(valid["going_min"].gt(previous_max).sum()),
        )
        results.append(row)
    result = pd.DataFrame(results)
    if result.empty:
        return pd.DataFrame(
            columns=key_columns
            + [
                "band_count", "inactive_band_count", "missing_bound_count",
                "invalid_bound_count", "overlap_count", "gap_count",
            ]
        )
    issue_columns = ["missing_bound_count", "invalid_bound_count", "overlap_count"]
    result["configuration_status"] = "valid"
    result.loc[result[issue_columns].sum(axis=1).gt(0), "configuration_status"] = "review"
    result.loc[
        result[issue_columns].sum(axis=1).eq(0) & result["gap_count"].gt(0),
        "configuration_status",
    ] = "valid_with_gaps"
    return result.sort_values(["course", "map_label"]).reset_index(drop=True)


def map_waypoint_compliance(
    going_reports: pd.DataFrame,
    stick_readings: pd.DataFrame,
    map_waypoints: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Compare each downloaded reading set with its map's configured waypoint values."""
    configured = map_waypoints.loc[
        map_waypoints["waypoint"].notna(), ["map_id", "waypoint"]
    ].drop_duplicates()
    configured["waypoint"] = configured["waypoint"].astype("int64")

    readings = stick_readings.copy()
    readings["sr_waypoint"] = pd.to_numeric(readings["sr_waypoint"], errors="coerce").astype("Int64")
    readings = readings.loc[readings["sr_waypoint"].notna()].copy()
    readings["sr_waypoint"] = readings["sr_waypoint"].astype("int64")
    observed = (
        readings.groupby(["sr_going_report_id", "gr_map_id", "sr_waypoint"], dropna=False)
        .agg(
            reading_rows=("sr_waypoint", "size"),
            valid_reading_rows=("data_type", lambda values: values.eq("data point").sum()),
        )
        .reset_index()
        .rename(
            columns={
                "sr_going_report_id": "gr_id",
                "gr_map_id": "map_id",
                "sr_waypoint": "waypoint",
            }
        )
    )
    report_maps = observed[["gr_id", "map_id"]].drop_duplicates()

    expected_detail = report_maps.merge(configured, on="map_id", how="inner").merge(
        observed,
        on=["gr_id", "map_id", "waypoint"],
        how="left",
    )
    expected_detail["waypoint_status"] = "expected_not_read"
    expected_detail.loc[
        expected_detail["reading_rows"].notna(), "waypoint_status"
    ] = "expected_and_read"

    observed_detail = observed.merge(
        configured.assign(configured_waypoint=True),
        on=["map_id", "waypoint"],
        how="left",
    )
    configured_maps = set(configured["map_id"].unique())
    unexpected_detail = observed_detail.loc[
        observed_detail["configured_waypoint"].ne(True)
    ].copy()
    unexpected_detail["waypoint_status"] = "unexpected_read"
    unexpected_detail.loc[
        ~unexpected_detail["map_id"].isin(configured_maps), "waypoint_status"
    ] = "no_map_configuration"
    detail_columns = [
        "gr_id", "map_id", "waypoint", "reading_rows", "valid_reading_rows",
        "waypoint_status",
    ]
    details = pd.concat(
        [expected_detail[detail_columns], unexpected_detail[detail_columns]],
        ignore_index=True,
    ).sort_values(["gr_id", "waypoint_status", "waypoint"]).reset_index(drop=True)

    expected_counts = configured.groupby("map_id")["waypoint"].nunique().rename(
        "expected_waypoint_count"
    )
    observed_counts = observed.groupby(["gr_id", "map_id"]).agg(
        observed_waypoint_count=("waypoint", "nunique"),
        valid_waypoint_count=("valid_reading_rows", lambda values: values.gt(0).sum()),
    )
    matched_counts = (
        expected_detail.loc[expected_detail["reading_rows"].notna()]
        .groupby(["gr_id", "map_id"])["waypoint"]
        .nunique()
        .rename("expected_waypoints_read")
    )
    unexpected_counts = (
        unexpected_detail.groupby(["gr_id", "map_id"])["waypoint"]
        .nunique()
        .rename("unexpected_waypoint_count")
    )
    summary = observed_counts.join(expected_counts, on="map_id").join(
        matched_counts, how="left"
    ).join(unexpected_counts, how="left").reset_index()
    count_columns = ["expected_waypoints_read", "unexpected_waypoint_count"]
    summary[count_columns] = summary[count_columns].fillna(0).astype("int64")
    summary["expected_waypoint_count"] = summary["expected_waypoint_count"].fillna(0).astype("int64")
    summary["missing_waypoint_count"] = (
        summary["expected_waypoint_count"] - summary["expected_waypoints_read"]
    ).clip(lower=0)
    summary["reading_count_matches"] = (
        summary["expected_waypoint_count"].gt(0)
        & summary["observed_waypoint_count"].eq(summary["expected_waypoint_count"])
    )
    summary["waypoint_set_matches"] = (
        summary["expected_waypoint_count"].gt(0)
        & summary["missing_waypoint_count"].eq(0)
        & summary["unexpected_waypoint_count"].eq(0)
    )
    summary["waypoint_compliance_status"] = "waypoint_set_matches"
    summary.loc[
        ~summary["waypoint_set_matches"], "waypoint_compliance_status"
    ] = "waypoint_set_mismatch"
    summary.loc[
        summary["expected_waypoint_count"].eq(0), "waypoint_compliance_status"
    ] = "no_map_configuration"

    report_columns = [
        column
        for column in ("gr_id", "race_date", "created_at_gr", "course", "map_id", "ma_label")
        if column in going_reports
    ]
    metadata = going_reports[report_columns].drop_duplicates("gr_id")
    summary = metadata.merge(summary, on=["gr_id", "map_id"], how="right")
    return (
        summary.sort_values(["race_date", "course", "gr_id"]).reset_index(drop=True),
        details,
    )


def categorise_stick_reading_averages(
    going_reports: pd.DataFrame,
    stick_readings: pd.DataFrame,
    going_groups: pd.DataFrame,
) -> pd.DataFrame:
    """Assign map-average stick indices to the map's configured Going bands."""
    valid_readings = stick_readings.copy()
    if "data_type" in valid_readings:
        valid_readings = valid_readings.loc[valid_readings["data_type"].eq("data point")]
    reading_summary = (
        valid_readings.groupby(["sr_going_report_id", "gr_map_id"], dropna=False)
        .agg(
            valid_reading_count=("sr_index", "count"),
            waypoint_count=("sr_waypoint", "nunique"),
            average_stick_index=("sr_index", "mean"),
            minimum_stick_index=("sr_index", "min"),
            maximum_stick_index=("sr_index", "max"),
        )
        .reset_index()
        .rename(columns={"sr_going_report_id": "gr_id", "gr_map_id": "map_id"})
    )

    report_columns = [
        "gr_id", "map_id", "goings_group_id", "official_going", "race_date",
        "created_at_gr", "course", "abandoned",
    ]
    reports = going_reports[report_columns].drop_duplicates("gr_id").copy()
    reports = reports.merge(reading_summary, on=["gr_id", "map_id"], how="left")

    bands = going_groups[
        ["map_id", "goings_group_id", "going_id", "going", "going_min", "going_max"]
    ].drop_duplicates()
    inactive = _normalise_text(bands["going"]).str.casefold().str.replace(" ", "").eq(
        "notinuse"
    )
    bands = bands.loc[
        ~inactive & bands["going_min"].notna() & bands["going_max"].notna()
    ].copy()
    bands["group_maximum"] = bands.groupby(
        ["map_id", "goings_group_id"], dropna=False
    )["going_max"].transform("max")
    candidates = reports.merge(
        bands,
        on=["map_id", "goings_group_id"],
        how="left",
    )
    candidates["band_match"] = (
        candidates["average_stick_index"].notna()
        & candidates["going_min"].notna()
        & candidates["going_max"].notna()
        & candidates["average_stick_index"].ge(candidates["going_min"])
        & (
            candidates["average_stick_index"].lt(candidates["going_max"])
            | (
                candidates["going_max"].eq(candidates["group_maximum"])
                & candidates["average_stick_index"].le(candidates["going_max"])
            )
        )
    )
    band_counts = (
        candidates.groupby("gr_id", dropna=False)["going_id"]
        .count()
        .rename("available_band_count")
    )
    matched = candidates.loc[candidates["band_match"]]
    assignments = (
        matched.groupby("gr_id", dropna=False)
        .agg(
            category_match_count=("going_id", "nunique"),
            calculated_going=("going", _joined_unique_values),
        )
    )
    result = reports.merge(band_counts, on="gr_id", how="left").merge(
        assignments, on="gr_id", how="left"
    )
    result["available_band_count"] = result["available_band_count"].fillna(0).astype("int64")
    result["category_match_count"] = result["category_match_count"].fillna(0).astype("int64")
    result["calculated_going"] = result["calculated_going"].fillna("")
    result["category_status"] = "assigned"
    result.loc[result["average_stick_index"].isna(), "category_status"] = "no_valid_readings"
    result.loc[
        result["average_stick_index"].notna() & result["goings_group_id"].isna(),
        "category_status",
    ] = "missing_going_group"
    result.loc[
        result["average_stick_index"].notna()
        & result["goings_group_id"].notna()
        & result["available_band_count"].eq(0),
        "category_status",
    ] = "no_group_configuration"
    result.loc[
        result["average_stick_index"].notna()
        & result["available_band_count"].gt(0)
        & result["category_match_count"].eq(0),
        "category_status",
    ] = "no_matching_band"
    result.loc[result["category_match_count"].gt(1), "category_status"] = "ambiguous_band_match"

    official = _normalise_text(result["official_going"]).str.casefold()
    calculated = _normalise_text(result["calculated_going"]).str.casefold()
    labels = sorted(
        {_normalise_text(pd.Series([label])).iloc[0].casefold() for label in bands["going"]},
        key=len,
        reverse=True,
    )
    label_pattern = re.compile(
        r"(?<!\w)(" + "|".join(re.escape(label) for label in labels) + r")(?!\w)"
    )
    reported_categories = official.map(lambda description: label_pattern.findall(description))
    result["reported_going_category"] = reported_categories.map(
        lambda categories: categories[0] if categories else ""
    )
    result["reported_going_categories"] = reported_categories.map(
        lambda categories: " | ".join(dict.fromkeys(categories))
    )
    result["going_match_status"] = "not_assigned"
    assigned = result["category_status"].eq("assigned")
    result.loc[assigned & official.eq(""), "going_match_status"] = "missing_official_going"
    primary_match = calculated.eq(result["reported_going_category"])
    result.loc[assigned & primary_match, "going_match_status"] = "primary_match"
    description_match = pd.Series(
        [
            bool(value) and value in categories
            for value, categories in zip(calculated, reported_categories)
        ],
        index=result.index,
    )
    result.loc[
        assigned & ~primary_match & description_match,
        "going_match_status",
    ] = "qualifier_match"
    result.loc[
        assigned & official.ne("") & ~primary_match & ~description_match,
        "going_match_status",
    ] = "mismatch"
    return result.sort_values(["race_date", "course", "gr_id"]).reset_index(drop=True)


def race_meeting_coverage(
    reconciled_calendar: pd.DataFrame,
    going_reports: pd.DataFrame,
    stick_readings: pd.DataFrame,
) -> pd.DataFrame:
    """Summarise whether each calendar race meeting has reports and stick readings."""
    optional_columns = [
        column
        for column in (
            "reconciliation_status", "database_record_count", "excel_record_count",
            "database_meeting_codes", "excel_race_codes",
        )
        if column in reconciled_calendar
    ]
    meetings = reconciled_calendar[["meeting_date", "course"] + optional_columns].drop_duplicates(
        ["meeting_date", "course"]
    ).copy()
    meetings["meeting_date"] = pd.to_datetime(meetings["meeting_date"]).dt.normalize()
    meetings["course"] = _normalise_course(meetings["course"])

    reports = going_reports.copy()
    reports["meeting_date"] = pd.to_datetime(reports["race_date"], errors="coerce").dt.normalize()
    reports["course"] = _normalise_course(reports["course"])
    meeting_reports = reports.merge(
        meetings[["meeting_date", "course"]],
        on=["meeting_date", "course"],
        how="inner",
    )

    report_summary = (
        meeting_reports.groupby(["meeting_date", "course"], dropna=False)
        .agg(
            report_count=("gr_id", "nunique"),
            first_report_created=("created_at_gr", "min"),
            last_report_created=("created_at_gr", "max"),
            abandoned_reports=("abandoned", "sum"),
        )
        .reset_index()
    )

    reading_summary = (
        stick_readings.drop(columns=["course"], errors="ignore").merge(
            meeting_reports[["gr_id", "meeting_date", "course"]].drop_duplicates(),
            left_on="sr_going_report_id",
            right_on="gr_id",
            how="inner",
        )
        .groupby(["meeting_date", "course"], dropna=False)
        .agg(
            stick_reading_count=("sr_going_report_id", "size"),
            valid_stick_reading_count=("data_type", lambda values: values.eq("data point").sum()),
        )
        .reset_index()
    )

    coverage = meetings.merge(report_summary, on=["meeting_date", "course"], how="left")
    coverage = coverage.merge(reading_summary, on=["meeting_date", "course"], how="left")
    count_columns = ["report_count", "abandoned_reports", "stick_reading_count", "valid_stick_reading_count"]
    coverage[count_columns] = coverage[count_columns].fillna(0).astype("int64")
    coverage["coverage_status"] = "report_with_stick_readings"
    coverage.loc[coverage["report_count"].eq(0), "coverage_status"] = "calendar_only"
    coverage.loc[
        coverage["report_count"].gt(0) & coverage["stick_reading_count"].eq(0),
        "coverage_status",
    ] = "report_without_stick_readings"
    coverage = coverage.sort_values(["meeting_date", "course"]).reset_index(drop=True)
    LOGGER.info("Race-meeting coverage complete: %s calendar meetings.", len(coverage))
    return coverage


def going_compliance_observations(
    reconciled_calendar: pd.DataFrame,
    going_reports: pd.DataFrame,
    stick_readings: pd.DataFrame,
    going_groups: pd.DataFrame,
    waypoint_compliance: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Evaluate Going evidence using calendar dates and report make_date."""
    key_columns = ["meeting_date", "course"]
    meetings = reconciled_calendar[key_columns].drop_duplicates().copy()
    meetings["meeting_date"] = pd.to_datetime(meetings["meeting_date"]).dt.normalize()
    meetings["course"] = _normalise_course(meetings["course"])
    meetings = meetings.loc[meetings["course"].ne("Wolverhampton")].copy()
    meetings["declaration_start_date"] = meetings["meeting_date"] - pd.Timedelta(days=3)
    meetings["declaration_date"] = meetings["meeting_date"] - pd.Timedelta(days=2)

    reports = going_reports.copy()
    if "stick_user" not in reports.columns:
        raise ValueError(
            "Going reports must include stick_user to exclude TT Auto Updater reports."
        )
    reports["stick_user_normalised"] = _normalise_text(reports["stick_user"]).str.casefold()
    reports["is_automated_official_going_source"] = reports[
        "stick_user_normalised"
    ].isin(AUTOMATED_OFFICIAL_GOING_USERS)
    excluded = reports["stick_user_normalised"].isin(EXCLUDED_GOING_REPORT_USERS)
    reports = reports.loc[~excluded].copy()
    LOGGER.info(
        "Retained %s automated API reports, including %s with Official Going.",
        int(reports["is_automated_official_going_source"].sum()),
        int(
            (
                reports["is_automated_official_going_source"]
                & _normalise_text(reports["official_going"]).ne("")
            ).sum()
        ),
    )
    reports["entered_race_date"] = pd.to_datetime(
        reports["race_date"], errors="coerce"
    ).dt.normalize()
    reports["created_date"] = pd.to_datetime(
        reports["created_at_gr"], errors="coerce"
    ).dt.normalize()
    reports["made_date"] = pd.to_datetime(
        reports["make_date"], errors="coerce"
    ).dt.normalize()
    reports["report_date"] = reports["made_date"].fillna(
        reports["created_date"]
    )
    reports["course"] = _normalise_course(reports["course"])

    # Match every report from its course and make_date against the racing
    # calendar, falling back to created_at only when make_date is absent. The
    # entered race_date is retained for audit and never controls classification.
    # A report may support more than one meeting window: for example, the
    # race-day going value for a multi-day meeting can also be the declaration
    # value for a later day.
    calendar_candidates = pd.concat(
        [
            meetings[key_columns].assign(
                report_date=meetings["meeting_date"],
                inferred_observation_window="race_day",
                candidate_priority=0,
            ),
            meetings[key_columns].assign(
                report_date=meetings["meeting_date"] - pd.Timedelta(days=2),
                inferred_observation_window="declaration_day",
                candidate_priority=1,
            ),
            meetings[key_columns].assign(
                report_date=meetings["meeting_date"] - pd.Timedelta(days=3),
                inferred_observation_window="declaration_day",
                candidate_priority=2,
            ),
        ],
        ignore_index=True,
    )
    meeting_reports = reports.merge(
        calendar_candidates,
        on=["course", "report_date"],
        how="inner",
    )
    meeting_reports = meeting_reports.drop(columns="candidate_priority")
    meeting_reports["meeting_date_source"] = "calendar_report_date"
    abandonment = (
        meeting_reports.groupby(key_columns, dropna=False)["abandoned"]
        .max()
        .fillna(0)
        .astype(bool)
        .rename("is_abandoned")
        .reset_index()
    )

    declaration = meetings.assign(
        observation_window="declaration_day",
        window_start_date=meetings["declaration_start_date"],
        required_date=meetings["declaration_date"],
    )
    race_day = meetings.assign(
        observation_window="race_day",
        window_start_date=meetings["meeting_date"],
        required_date=meetings["meeting_date"],
    )
    observations = pd.concat([declaration, race_day], ignore_index=True)
    observations = observations.merge(abandonment, on=key_columns, how="left")
    observations["is_abandoned"] = observations["is_abandoned"].fillna(False).astype(bool)

    eligible_reports = observations[
        key_columns + ["observation_window", "window_start_date", "required_date"]
    ].merge(meeting_reports, on=key_columns, how="inner")
    eligible_reports = eligible_reports.loc[
        eligible_reports["report_date"].between(
            eligible_reports["window_start_date"],
            eligible_reports["required_date"],
        )
    ].copy()

    # Once a report is assigned to a calendar window by make_date (falling
    # back to created_at), all readings linked to that report belong to the
    # same window. Their individual sr_created_at timestamps do not reassign
    # or exclude them.
    readings = stick_readings.copy()
    report_windows = eligible_reports[["gr_id"]].drop_duplicates().rename(
        columns={"gr_id": "sr_going_report_id"}
    )
    eligible_readings = readings.merge(
        report_windows,
        on="sr_going_report_id",
        how="inner",
    )

    categorised = categorise_stick_reading_averages(
        eligible_reports,
        eligible_readings,
        going_groups,
    )
    evidence_keys = eligible_reports[
        ["gr_id"]
        + key_columns
        + ["observation_window", "window_start_date", "required_date"]
    ].drop_duplicates()
    evidence = categorised.drop(columns=["course"], errors="ignore").merge(
        evidence_keys, on="gr_id", how="left"
    )
    if waypoint_compliance is None:
        evidence["waypoint_set_matches"] = True
    else:
        evidence = evidence.merge(
            waypoint_compliance[["gr_id", "waypoint_set_matches"]].drop_duplicates("gr_id"),
            on="gr_id",
            how="left",
        )
        evidence["waypoint_set_matches"] = evidence["waypoint_set_matches"].fillna(False)
    evidence["going_reported"] = _normalise_text(evidence["official_going"]).ne("")
    evidence["going_matches"] = (
        evidence["going_match_status"].isin(["primary_match", "qualifier_match"])
        & evidence["waypoint_set_matches"]
    )
    evidence["going_mismatch"] = evidence["going_match_status"].eq("mismatch")

    summary_keys = key_columns + [
        "observation_window",
        "window_start_date",
        "required_date",
    ]
    if evidence.empty:
        evidence_summary = pd.DataFrame(columns=summary_keys)
    else:
        evidence_summary = (
            evidence.groupby(summary_keys, dropna=False)
            .agg(
                report_count=("gr_id", "nunique"),
                going_reported_count=("going_reported", "sum"),
                valid_stick_reading_count=("valid_reading_count", "sum"),
                complete_waypoint_set_count=("waypoint_set_matches", "sum"),
                matching_report_count=("going_matches", "sum"),
                mismatch_report_count=("going_mismatch", "sum"),
                evidence_report_ids=("gr_id", lambda values: " | ".join(map(str, sorted(values.unique())))),
            )
            .reset_index()
        )
    result = observations.merge(evidence_summary, on=summary_keys, how="left")
    count_columns = [
        "report_count", "going_reported_count", "valid_stick_reading_count",
        "complete_waypoint_set_count", "matching_report_count", "mismatch_report_count",
    ]
    result[count_columns] = result[count_columns].fillna(0).astype("int64")
    result["evidence_report_ids"] = result["evidence_report_ids"].fillna("")
    result["window_status"] = "going_matches"
    result.loc[result["matching_report_count"].eq(0), "window_status"] = "going_mismatch"
    result.loc[
        result["valid_stick_reading_count"].gt(0)
        & result["matching_report_count"].eq(0)
        & result["mismatch_report_count"].eq(0),
        "window_status",
    ] = "categorisation_review"
    result.loc[
        result["valid_stick_reading_count"].gt(0)
        & result["complete_waypoint_set_count"].eq(0),
        "window_status",
    ] = "waypoint_set_mismatch"
    result.loc[
        result["report_count"].gt(0) & result["valid_stick_reading_count"].eq(0),
        "window_status",
    ] = "stick_readings_not_downloaded"
    result.loc[
        result["report_count"].gt(0) & result["going_reported_count"].eq(0),
        "window_status",
    ] = "going_not_reported"
    result.loc[result["report_count"].eq(0), "window_status"] = "missing_report"
    result.loc[result["is_abandoned"], "window_status"] = "abandoned_excluded"

    result["reporting_category"] = "data uploaded - Going matches"
    result.loc[
        result["window_status"].isin(
            ["going_mismatch", "categorisation_review", "waypoint_set_mismatch"]
        ),
        "reporting_category",
    ] = "data uploaded - mismatch or review"
    result.loc[
        result["window_status"].eq("stick_readings_not_downloaded"),
        "reporting_category",
    ] = "going value reported"
    result.loc[
        result["window_status"].isin(["missing_report", "going_not_reported"]),
        "reporting_category",
    ] = "going value not reported"
    result.loc[result["is_abandoned"], "reporting_category"] = "abandoned"
    result["year"] = result["meeting_date"].dt.year
    return result.sort_values(key_columns + ["observation_window"]).reset_index(drop=True)


def meeting_compliance(observations: pd.DataFrame) -> pd.DataFrame:
    """Collapse the two required observations into one meeting-level decision."""
    key_columns = ["meeting_date", "course"]
    statuses = observations.pivot(
        index=key_columns,
        columns="observation_window",
        values="window_status",
    ).reset_index()
    statuses.columns.name = None
    abandoned = (
        observations.groupby(key_columns, dropna=False)["is_abandoned"]
        .max()
        .reset_index()
    )
    result = statuses.merge(abandoned, on=key_columns, how="left")
    result["declaration_compliant"] = result["declaration_day"].eq("going_matches")
    result["race_day_compliant"] = result["race_day"].eq("going_matches")
    result["compliance_status"] = "non_compliant"
    result.loc[
        result["declaration_compliant"] & result["race_day_compliant"],
        "compliance_status",
    ] = "compliant"
    result.loc[result["is_abandoned"], "compliance_status"] = "abandoned_excluded"
    result["combined_reporting_category"] = "neither day compliant"
    result.loc[
        result["declaration_compliant"] & ~result["race_day_compliant"],
        "combined_reporting_category",
    ] = "declaration day only"
    result.loc[
        ~result["declaration_compliant"] & result["race_day_compliant"],
        "combined_reporting_category",
    ] = "race day only"
    result.loc[
        result["declaration_compliant"] & result["race_day_compliant"],
        "combined_reporting_category",
    ] = "both days compliant"
    result.loc[result["is_abandoned"], "combined_reporting_category"] = "abandoned"
    result["year"] = pd.to_datetime(result["meeting_date"]).dt.year
    return result.sort_values(key_columns).reset_index(drop=True)


def non_meeting_activity(
    calendar_database: pd.DataFrame,
    going_reports: pd.DataFrame,
    stick_readings: pd.DataFrame,
) -> pd.DataFrame:
    """Summarise reports/readings that are not linked to an ASCOT calendar meeting."""
    meeting_keys = calendar_database[["meeting_date", "course"]].drop_duplicates().copy()
    meeting_keys["meeting_date"] = pd.to_datetime(meeting_keys["meeting_date"]).dt.normalize()
    meeting_keys["course"] = _normalise_course(meeting_keys["course"])

    reports = going_reports.copy()
    reports["meeting_date"] = pd.to_datetime(reports["race_date"], errors="coerce").dt.normalize()
    reports["activity_date"] = pd.to_datetime(reports["created_at_gr"]).dt.normalize()
    reports["course"] = _normalise_course(reports["course"])
    reports = reports.merge(
        meeting_keys.assign(on_calendar=True),
        on=["meeting_date", "course"],
        how="left",
    )
    non_meeting_reports = reports.loc[reports["on_calendar"].ne(True)].copy()

    non_meeting_readings = stick_readings.drop(columns=["course"], errors="ignore").merge(
        non_meeting_reports[["gr_id", "activity_date", "course"]],
        left_on="sr_going_report_id",
        right_on="gr_id",
        how="inner",
    )
    summary = (
        non_meeting_reports.groupby(["activity_date", "course"], dropna=False)
        .agg(non_meeting_reports=("gr_id", "nunique"))
        .reset_index()
        .merge(
            non_meeting_readings.groupby(["activity_date", "course"], dropna=False)
            .agg(
                non_meeting_stick_readings=("sr_going_report_id", "size"),
                valid_non_meeting_stick_readings=("data_type", lambda values: values.eq("data point").sum()),
            )
            .reset_index(),
            on=["activity_date", "course"],
            how="left",
        )
    )
    count_columns = [
        "non_meeting_reports",
        "non_meeting_stick_readings",
        "valid_non_meeting_stick_readings",
    ]
    summary[count_columns] = summary[count_columns].fillna(0).astype("int64")
    summary = summary.sort_values(["activity_date", "course"]).reset_index(drop=True)
    LOGGER.info("Non-meeting activity summary complete: %s date-course rows.", len(summary))
    return summary
